#!/usr/bin/env python3
"""
대시보드 · 판매후보 분석 빌더
=============================

scrape.py 가 만든 output/data.json 을 읽어 산출물을 생성한다.
목표: "직접 디자인해서 팔 만한, 원가(단가)가 낮은 제품" 을 골라내기 위한 분석.

산출물
    1) 엑셀 (.xlsx)
        · "저단가 판매후보" 시트 — 배송비까지 조회된 후보를 원가 낮은 순으로.
          이미지 + 원가/권장정가/마진/마진율 + 국가별 최저배송 + 예상 착지원가
        · "전체 카탈로그(단가순)" 시트 — 수집된 모든 제품을 원가순으로(이미지 없음, 마스터)
        · "배송비 상세" 시트 — (후보 × 국가 × 배송사) 롱포맷
        · "README" 시트
    2) HTML 대시보드 (후보)
    3) CSV — 전체 카탈로그 단가순 (엑셀/구글시트에서 자유 필터)

용어
    · 원가(base)    = 사이트 판매가(products-list-price). POD 공급가 ≒ 내가 물건에 지불하는 값.
    · 권장정가(retail) = 사이트 Retail Price(제안 소비자가).
    · 마진(margin)   = 권장정가 - 원가.  마진율 = 마진 / 권장정가.
    낮은 원가 = 내 디자인 마진을 붙일 여유가 큼 → 판매 후보로서 매력적.
"""
from __future__ import annotations

import argparse
import csv
import html as html_lib
import json
from collections import defaultdict
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OUTPUT_DIR = HERE / "output"
IMG_CACHE = OUTPUT_DIR / "images"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120.0 Safari/537.36")

THUMB = 90
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
GOOD_FILL = PatternFill("solid", fgColor="DCFCE7")  # 높은 마진율 강조


def load(data_path: Path) -> dict:
    return json.loads(data_path.read_text(encoding="utf-8"))


def download_image(url: str, pid: str) -> Path | None:
    if not url:
        return None
    if url.startswith("//"):
        url = "https:" + url
    IMG_CACHE.mkdir(parents=True, exist_ok=True)
    dest = IMG_CACHE / f"{pid}.jpg"
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    try:
        r = requests.get(url, headers={"User-Agent": UA}, timeout=30)
        if r.status_code == 200 and r.content:
            dest.write_bytes(r.content)
            return dest
    except requests.RequestException:
        pass
    return None


def shipping_by_country(product: dict) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for opt in product.get("shipping", []):
        grouped[opt["country"]].append(opt)
    return grouped


def min_price(opts: list[dict]) -> float | None:
    vals = [o["price"] for o in opts if o.get("price") is not None]
    return min(vals) if vals else None


def margin(product: dict) -> tuple[float | None, float | None]:
    base = product.get("sale_price")
    retail = product.get("retail_price")
    if base is None or retail is None:
        return None, None
    m = retail - base
    pct = (m / retail * 100) if retail else None
    return m, pct


def usd(cell, val):
    cell.value = val
    cell.number_format = '"$"#,##0.00'


# --------------------------------------------------------------------------- #
# 엑셀
# --------------------------------------------------------------------------- #
def build_excel(data: dict, xlsx_path: Path) -> None:
    products = data["products"]
    countries = [c["name"] for c in data["countries"]]
    candidates = [p for p in products if p.get("shipping")]
    candidates.sort(key=lambda p: (p.get("sale_price") is None, p.get("sale_price") or 0))
    catalog = sorted(products, key=lambda p: (p.get("sale_price") is None,
                                              p.get("sale_price") or 0))

    wb = Workbook()

    # ---------- 시트 1: 저단가 판매후보 (이미지 + 배송비) ---------- #
    ws = wb.active
    ws.title = "저단가 판매후보"
    headers = (["이미지", "제품ID", "제품명", "카테고리", "원가(USD)", "권장정가",
                "마진", "마진율%"]
               + [f"{c} 최저배송" for c in countries]
               + ["예상착지원가(원가+US최저배송)", "제품 링크"])
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = CENTER; c.border = BORDER
    ws.freeze_panes = "C2"
    widths = [14, 8, 40, 18, 11, 11, 10, 9] + [13] * len(countries) + [22, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    row = 2
    for p in candidates:
        ws.row_dimensions[row].height = THUMB * 0.78
        img_path = download_image(p.get("image", ""), p["pid"])
        if img_path:
            try:
                xi = XLImage(str(img_path)); xi.width = xi.height = THUMB
                ws.add_image(xi, f"A{row}")
            except Exception:  # noqa: BLE001
                pass
        ws.cell(row=row, column=2, value=p.get("pid", ""))
        ws.cell(row=row, column=3, value=p.get("name", ""))
        ws.cell(row=row, column=4, value=p.get("category_name", ""))
        usd(ws.cell(row=row, column=5), p.get("sale_price"))
        usd(ws.cell(row=row, column=6), p.get("retail_price"))
        m, pct = margin(p)
        usd(ws.cell(row=row, column=7), m)
        pctcell = ws.cell(row=row, column=8, value=(round(pct, 1) if pct is not None else None))
        grouped = shipping_by_country(p)
        col = 9
        us_min = None
        for cname in countries:
            mp = min_price(grouped.get(cname, []))
            usd(ws.cell(row=row, column=col), mp)
            if cname.startswith("미국"):
                us_min = mp
            col += 1
        landed = None
        if p.get("sale_price") is not None and us_min is not None:
            landed = p["sale_price"] + us_min
        usd(ws.cell(row=row, column=col), landed)
        link = ws.cell(row=row, column=col + 1, value=p.get("url", ""))
        link.hyperlink = p.get("url", ""); link.font = Font(color="2563EB", underline="single")

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=c)
            cell.border = BORDER
            cell.alignment = LEFT if c in (3, 4) else CENTER
        if pct is not None and pct >= 60:
            ws.cell(row=row, column=8).fill = GOOD_FILL
        row += 1

    # ---------- 시트 2: 전체 카탈로그(단가순) — 이미지 없음 ---------- #
    ws2 = wb.create_sheet("전체 카탈로그(단가순)")
    h2 = ["순위", "제품ID", "제품명", "카테고리", "원가(USD)", "권장정가",
          "마진", "마진율%", "이미지 URL", "제품 링크"]
    for col, h in enumerate(h2, start=1):
        c = ws2.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = CENTER; c.border = BORDER
    ws2.freeze_panes = "A2"
    for i, w in enumerate([6, 8, 46, 20, 11, 11, 10, 9, 40, 26], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for rank, p in enumerate(catalog, start=1):
        ws2.cell(row=r, column=1, value=rank)
        ws2.cell(row=r, column=2, value=p.get("pid", ""))
        ws2.cell(row=r, column=3, value=p.get("name", ""))
        ws2.cell(row=r, column=4, value=p.get("category_name", ""))
        usd(ws2.cell(row=r, column=5), p.get("sale_price"))
        usd(ws2.cell(row=r, column=6), p.get("retail_price"))
        m, pct = margin(p)
        usd(ws2.cell(row=r, column=7), m)
        ws2.cell(row=r, column=8, value=(round(pct, 1) if pct is not None else None))
        ws2.cell(row=r, column=9, value=p.get("image", ""))
        lk = ws2.cell(row=r, column=10, value=p.get("url", ""))
        lk.hyperlink = p.get("url", ""); lk.font = Font(color="2563EB", underline="single")
        for c in range(1, 11):
            ws2.cell(row=r, column=c).alignment = LEFT if c in (3, 4, 9) else CENTER
        r += 1
    ws2.auto_filter.ref = f"A1:J{max(1, r - 1)}"

    # ---------- 시트 3: 배송비 상세(후보) ---------- #
    ws3 = wb.create_sheet("배송비 상세")
    h3 = ["제품ID", "제품명", "카테고리", "원가(USD)", "국가", "국가코드",
          "배송사/서비스", "배송비(USD)", "제품 링크"]
    for col, h in enumerate(h3, start=1):
        c = ws3.cell(row=1, column=col, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.alignment = CENTER; c.border = BORDER
    ws3.freeze_panes = "A2"
    for i, w in enumerate([8, 40, 18, 11, 12, 10, 28, 13, 26], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    r = 2
    for p in candidates:
        for o in p.get("shipping", []):
            ws3.cell(row=r, column=1, value=p.get("pid", ""))
            ws3.cell(row=r, column=2, value=p.get("name", ""))
            ws3.cell(row=r, column=3, value=p.get("category_name", ""))
            usd(ws3.cell(row=r, column=4), p.get("sale_price"))
            ws3.cell(row=r, column=5, value=o.get("country", ""))
            ws3.cell(row=r, column=6, value=o.get("country_code", ""))
            ws3.cell(row=r, column=7, value=o.get("carrier", ""))
            usd(ws3.cell(row=r, column=8), o.get("price"))
            ws3.cell(row=r, column=9, value=p.get("url", ""))
            for c in range(1, 10):
                ws3.cell(row=r, column=c).alignment = LEFT if c in (2, 3, 7) else CENTER
            r += 1
    ws3.auto_filter.ref = f"A1:I{max(1, r - 1)}"

    # ---------- 시트 4: README ---------- #
    ws4 = wb.create_sheet("README")
    prices = [p["sale_price"] for p in products if p.get("sale_price") is not None]
    notes = [
        ("InterestPrint 판매후보 분석", True),
        (f"수집 시각: {data.get('scraped_at', '')}", False),
        (f"전체 수집 제품: {len(products)}개  |  배송비 조회 후보: {len(candidates)}개", False),
        (f"원가 범위: ${min(prices):.2f} ~ ${max(prices):.2f}  (제품 {len(prices)}개 기준)" if prices else "", False),
        (f"조회 국가: {', '.join(countries)}  (수량 1개 기준)", False),
        ("", False),
        ("· '저단가 판매후보': 배송비까지 조회된 제품을 원가 낮은 순으로. 마진율 60%+ 는 초록 강조", False),
        ("· '전체 카탈로그(단가순)': 수집된 모든 제품의 원가/마진 마스터 목록(자동필터)", False),
        ("· '배송비 상세': (제품×국가×배송사) 롱포맷 — 피벗/필터 분석용", False),
        ("· 원가=사이트 판매가(POD 공급가), 권장정가=Retail Price, 마진=정가-원가", False),
        ("· 낮은 원가일수록 내 디자인 마진 여유가 큼. 배송비까지 더한 '예상착지원가' 참고", False),
        ("· 배송비는 수량 1개 기준 실시간 API 값(shipping_pirce_estimate)", False),
    ]
    for i, (line, bold) in enumerate(notes, start=1):
        c = ws4.cell(row=i, column=1, value=line)
        if bold:
            c.font = Font(bold=True, size=14)
    ws4.column_dimensions["A"].width = 95

    # README 시트를 맨 앞이 아닌 뒤로 두고, 저단가 후보를 활성 시트로
    wb.active = 0
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx_path)
    print(f"✓ 엑셀 저장 → {xlsx_path}  (후보 {len(candidates)} / 전체 {len(products)})")


# --------------------------------------------------------------------------- #
# CSV (전체 카탈로그 단가순)
# --------------------------------------------------------------------------- #
def build_csv(data: dict, csv_path: Path) -> None:
    products = sorted(data["products"],
                      key=lambda p: (p.get("sale_price") is None, p.get("sale_price") or 0))
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["순위", "제품ID", "제품명", "카테고리", "원가USD", "권장정가USD",
                    "마진USD", "마진율%", "배송비조회됨", "이미지URL", "제품링크"])
        for rank, p in enumerate(products, start=1):
            m, pct = margin(p)
            w.writerow([rank, p.get("pid", ""), p.get("name", ""), p.get("category_name", ""),
                        p.get("sale_price", ""), p.get("retail_price", ""),
                        round(m, 2) if m is not None else "",
                        round(pct, 1) if pct is not None else "",
                        "Y" if p.get("shipping") else "",
                        p.get("image", ""), p.get("url", "")])
    print(f"✓ CSV 저장 → {csv_path}")


# --------------------------------------------------------------------------- #
# HTML (후보)
# --------------------------------------------------------------------------- #
def build_html(data: dict, html_path: Path) -> None:
    countries = [c["name"] for c in data["countries"]]
    candidates = [p for p in data["products"] if p.get("shipping")]
    candidates.sort(key=lambda p: (p.get("sale_price") is None, p.get("sale_price") or 0))
    esc = html_lib.escape

    cards = []
    for p in candidates:
        grouped = shipping_by_country(p)
        ship_html = []
        for cname in countries:
            opts = grouped.get(cname, [])
            mp = min_price(opts)
            rows = "".join(
                f'<tr><td>{esc(o["carrier"])}</td>'
                f'<td class="pr">{esc(o["price_text"] or "-")}</td></tr>'
                for o in opts) or '<tr><td colspan="2">배송 옵션 없음</td></tr>'
            mp_txt = f"${mp:,.2f}" if mp is not None else "-"
            ship_html.append(f'<div class="country"><div class="chead"><span>{esc(cname)}</span>'
                             f'<span class="min">최저 {mp_txt}</span></div>'
                             f'<table class="stbl"><tbody>{rows}</tbody></table></div>')
        m, pct = margin(p)
        sale = f'${p["sale_price"]:,.2f}' if p.get("sale_price") is not None else "-"
        retail = f'${p["retail_price"]:,.2f}' if p.get("retail_price") is not None else ""
        margin_txt = (f'<span class="mg">마진 ${m:,.2f} ({pct:.0f}%)</span>'
                      if m is not None else "")
        img = p.get("image", "")
        if img.startswith("//"):
            img = "https:" + img
        cards.append(f"""
        <div class="card">
          <a href="{esc(p.get('url',''))}" target="_blank" rel="noopener">
            <img loading="lazy" src="{esc(img)}" alt="{esc(p.get('name',''))}"></a>
          <div class="body">
            <div class="cat">{esc(p.get('category_name',''))}</div>
            <a class="name" href="{esc(p.get('url',''))}" target="_blank" rel="noopener">{esc(p.get('name',''))}</a>
            <div class="price"><span class="sale">{sale}</span>
              {'<span class="retail">'+retail+'</span>' if retail else ''}{margin_txt}</div>
            <div class="ship">{''.join(ship_html)}</div>
          </div>
        </div>""")

    doc = f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>InterestPrint 저단가 판매후보</title>
<style>
  :root {{ color-scheme: light dark; }}
  * {{ box-sizing: border-box; }}
  body {{ margin:0; font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Malgun Gothic",sans-serif;
         background:#f6f7f9; color:#1f2937; }}
  header {{ background:#111827; color:#fff; padding:20px 24px; }}
  header h1 {{ margin:0 0 4px; font-size:20px; }}
  header p {{ margin:0; color:#9ca3af; font-size:13px; }}
  .grid {{ display:grid; grid-template-columns:repeat(auto-fill,minmax(300px,1fr));
           gap:16px; padding:20px; max-width:1400px; margin:0 auto; }}
  .card {{ background:#fff; border:1px solid #e5e7eb; border-radius:12px; overflow:hidden;
           display:flex; flex-direction:column; }}
  .card > a {{ display:block; background:#fff; text-align:center; }}
  .card img {{ width:100%; max-width:220px; height:auto; margin:12px auto 0; }}
  .body {{ padding:12px 14px 16px; display:flex; flex-direction:column; gap:8px; }}
  .cat {{ font-size:11px; color:#6b7280; text-transform:uppercase; letter-spacing:.03em; }}
  .name {{ font-weight:600; font-size:14px; line-height:1.35; color:#111827; text-decoration:none; }}
  .name:hover {{ color:#2563eb; }}
  .price {{ display:flex; align-items:baseline; gap:8px; flex-wrap:wrap; }}
  .sale {{ font-size:18px; font-weight:700; color:#dc2626; }}
  .retail {{ font-size:13px; color:#9ca3af; text-decoration:line-through; }}
  .mg {{ font-size:12px; color:#059669; font-weight:600; }}
  .ship {{ display:flex; flex-direction:column; gap:8px; margin-top:4px; }}
  .country {{ border:1px solid #eef0f3; border-radius:8px; overflow:hidden; }}
  .chead {{ display:flex; justify-content:space-between; align-items:center;
            background:#eef2ff; padding:6px 10px; font-size:12px; font-weight:600; }}
  .chead .min {{ color:#4338ca; }}
  .stbl {{ width:100%; border-collapse:collapse; font-size:12px; }}
  .stbl td {{ padding:4px 10px; border-top:1px solid #f1f2f4; }}
  .stbl td.pr {{ text-align:right; font-variant-numeric:tabular-nums; white-space:nowrap; }}
  @media (prefers-color-scheme: dark) {{
    body {{ background:#0b0f17; color:#e5e7eb; }}
    .card {{ background:#111827; border-color:#1f2937; }}
    .card > a {{ background:#fff; }}
    .name {{ color:#f3f4f6; }}
    .country {{ border-color:#1f2937; }}
    .chead {{ background:#1e293b; }}
    .chead .min {{ color:#a5b4fc; }}
    .stbl td {{ border-color:#1f2937; }}
  }}
</style></head><body>
<header>
  <h1>InterestPrint 저단가 판매후보 (원가 낮은 순)</h1>
  <p>후보 {len(candidates)}개 · 조회 국가: {esc(', '.join(countries))} (수량 1개 기준) · 수집 {esc(data.get('scraped_at',''))}</p>
</header>
<div class="grid">{''.join(cards)}</div>
</body></html>"""
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(doc, encoding="utf-8")
    print(f"✓ HTML 저장 → {html_path}")


def main() -> None:
    ap = argparse.ArgumentParser(description="InterestPrint 판매후보 분석 빌더")
    ap.add_argument("--data", default=str(OUTPUT_DIR / "data.json"))
    ap.add_argument("--xlsx", default=str(OUTPUT_DIR / "interestprint_dashboard.xlsx"))
    ap.add_argument("--html", default=str(OUTPUT_DIR / "dashboard.html"))
    ap.add_argument("--csv", default=str(OUTPUT_DIR / "catalog_by_price.csv"))
    args = ap.parse_args()

    data = load(Path(args.data))
    build_excel(data, Path(args.xlsx))
    build_csv(data, Path(args.csv))
    build_html(data, Path(args.html))


if __name__ == "__main__":
    main()
